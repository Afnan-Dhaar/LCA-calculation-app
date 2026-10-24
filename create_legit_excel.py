"""
Script to generate a production-grade, beautifully formatted Excel workbook
containing real-world Life Cycle Assessment (LCA) product models.
"""

from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PRODUCTS = [
    {
        "product_id": "flask_750_steel",
        "product_name": "Insulated Stainless Steel Water Bottle (750ml)",
        "functional_unit_amount": 1.0,
        "functional_unit": "unit (1 reusable bottle)",
    },
    {
        "product_id": "office_task_chair",
        "product_name": "Ergonomic Office Task Chair",
        "functional_unit_amount": 1.0,
        "functional_unit": "unit (1 ergonomic chair)",
    },
]

PROCESSES = [
    # -------------------------------------------------------------
    # 1. Insulated Stainless Steel Flask (Cradle-to-Grave)
    # -------------------------------------------------------------
    {
        "process_id": "flask_raw_materials",
        "product_id": "flask_750_steel",
        "process_name": "Raw Material Extraction & Compounding",
        "stage": "Raw Materials",
    },
    {
        "process_id": "flask_manufacturing",
        "product_id": "flask_750_steel",
        "process_name": "Deep-Drawing, Vacuum Brazing & Flask Assembly",
        "stage": "Manufacturing",
    },
    {
        "process_id": "flask_distribution",
        "product_id": "flask_750_steel",
        "process_name": "Global Maritime & Regional Retail Logistics",
        "stage": "Distribution",
    },
    {
        "process_id": "flask_end_of_life",
        "product_id": "flask_750_steel",
        "process_name": "End-of-Life Sorting & Metal Scrap Recycling",
        "stage": "End of Life",
    },
    # -------------------------------------------------------------
    # 2. Ergonomic Office Task Chair (Cradle-to-Gate + Distribution)
    # -------------------------------------------------------------
    {
        "process_id": "chair_materials",
        "product_id": "office_task_chair",
        "process_name": "Structural Metals & Polymer Injection Molding",
        "stage": "Raw Materials",
    },
    {
        "process_id": "chair_assembly",
        "product_id": "office_task_chair",
        "process_name": "Robotic Mechanical Assembly & Packaging",
        "stage": "Manufacturing",
    },
    {
        "process_id": "chair_logistics",
        "product_id": "office_task_chair",
        "process_name": "Commercial Freight Distribution",
        "stage": "Distribution",
    },
]

EXCHANGES = [
    # -------------------------------------------------------------
    # Flask: Raw Materials
    # -------------------------------------------------------------
    {
        "process_id": "flask_raw_materials",
        "flow_name": "Steel, structural generic",
        "exchange_type": "input",
        "amount": 0.280,
        "unit": "kg",
    },
    {
        "process_id": "flask_raw_materials",
        "flow_name": "Plastic granules / pellets (generic)",
        "exchange_type": "input",
        "amount": 0.045,
        "unit": "kg",
    },
    {
        "process_id": "flask_raw_materials",
        "flow_name": "Copper, refined generic",
        "exchange_type": "input",
        "amount": 0.008,
        "unit": "kg",
    },
    {
        "process_id": "flask_raw_materials",
        "flow_name": "Cardboard, corrugated packaging",
        "exchange_type": "input",
        "amount": 0.065,
        "unit": "kg",
    },
    # -------------------------------------------------------------
    # Flask: Manufacturing & Assembly
    # -------------------------------------------------------------
    {
        "process_id": "flask_manufacturing",
        "flow_name": "Raw Material Extraction & Compounding",
        "exchange_type": "input",
        "amount": 1.0,
        "unit": "unit",
    },
    {
        "process_id": "flask_manufacturing",
        "flow_name": "Electricity, grid average",
        "exchange_type": "input",
        "amount": 1.150,
        "unit": "kWh",
    },
    {
        "process_id": "flask_manufacturing",
        "flow_name": "Process heat, from natural gas boiler",
        "exchange_type": "input",
        "amount": 0.450,
        "unit": "kWh",
    },
    {
        "process_id": "flask_manufacturing",
        "flow_name": "Tap water, municipal",
        "exchange_type": "input",
        "amount": 2.800,
        "unit": "kg",
    },
    {
        "process_id": "flask_manufacturing",
        "flow_name": "Carbon dioxide, fossil",
        "exchange_type": "output",
        "amount": 0.042,
        "unit": "kg",
    },
    {
        "process_id": "flask_manufacturing",
        "flow_name": "Nitrogen oxides",
        "exchange_type": "output",
        "amount": 0.0003,
        "unit": "kg",
    },
    # -------------------------------------------------------------
    # Flask: Distribution
    # -------------------------------------------------------------
    {
        "process_id": "flask_distribution",
        "flow_name": "Deep-Drawing, Vacuum Brazing & Flask Assembly",
        "exchange_type": "input",
        "amount": 1.0,
        "unit": "unit",
    },
    # -------------------------------------------------------------
    # Flask: End of Life
    # -------------------------------------------------------------
    {
        "process_id": "flask_end_of_life",
        "flow_name": "Global Maritime & Regional Retail Logistics",
        "exchange_type": "input",
        "amount": 1.0,
        "unit": "unit",
    },
    {
        "process_id": "flask_end_of_life",
        "flow_name": "Electricity, grid average",
        "exchange_type": "input",
        "amount": 0.120,
        "unit": "kWh",
    },
    {
        "process_id": "flask_end_of_life",
        "flow_name": "Carbon dioxide, fossil",
        "exchange_type": "output",
        "amount": 0.015,
        "unit": "kg",
    },
    # -------------------------------------------------------------
    # Chair: Raw Materials
    # -------------------------------------------------------------
    {
        "process_id": "chair_materials",
        "flow_name": "Steel, structural generic",
        "exchange_type": "input",
        "amount": 5.800,
        "unit": "kg",
    },
    {
        "process_id": "chair_materials",
        "flow_name": "Aluminum, primary/secondary generic",
        "exchange_type": "input",
        "amount": 2.400,
        "unit": "kg",
    },
    {
        "process_id": "chair_materials",
        "flow_name": "Plastic granules / pellets (generic)",
        "exchange_type": "input",
        "amount": 4.200,
        "unit": "kg",
    },
    {
        "process_id": "chair_materials",
        "flow_name": "Cardboard, corrugated packaging",
        "exchange_type": "input",
        "amount": 1.800,
        "unit": "kg",
    },
    # -------------------------------------------------------------
    # Chair: Assembly & Packaging
    # -------------------------------------------------------------
    {
        "process_id": "chair_assembly",
        "flow_name": "Structural Metals & Polymer Injection Molding",
        "exchange_type": "input",
        "amount": 1.0,
        "unit": "unit",
    },
    {
        "process_id": "chair_assembly",
        "flow_name": "Electricity, grid average",
        "exchange_type": "input",
        "amount": 8.500,
        "unit": "kWh",
    },
    {
        "process_id": "chair_assembly",
        "flow_name": "Carbon dioxide, fossil",
        "exchange_type": "output",
        "amount": 0.350,
        "unit": "kg",
    },
    # -------------------------------------------------------------
    # Chair: Logistics
    # -------------------------------------------------------------
    {
        "process_id": "chair_logistics",
        "flow_name": "Robotic Mechanical Assembly & Packaging",
        "exchange_type": "input",
        "amount": 1.0,
        "unit": "unit",
    },
]

TRANSPORTATION = [
    # -------------------------------------------------------------
    # Flask: Transport Legs
    # -------------------------------------------------------------
    {
        "process_id": "flask_raw_materials",
        "transport_mode": "train",
        "distance": 450.0,
        "distance_unit": "km",
        "transported_mass": 0.400,
        "mass_unit": "kg",
    },
    {
        "process_id": "flask_manufacturing",
        "transport_mode": "truck",
        "distance": 120.0,
        "distance_unit": "km",
        "transported_mass": 0.400,
        "mass_unit": "kg",
    },
    {
        "process_id": "flask_distribution",
        "transport_mode": "ship",
        "distance": 2400.0,
        "distance_unit": "km",
        "transported_mass": 0.400,
        "mass_unit": "kg",
    },
    {
        "process_id": "flask_distribution",
        "transport_mode": "truck",
        "distance": 180.0,
        "distance_unit": "km",
        "transported_mass": 0.400,
        "mass_unit": "kg",
    },
    {
        "process_id": "flask_end_of_life",
        "transport_mode": "truck",
        "distance": 50.0,
        "distance_unit": "km",
        "transported_mass": 0.350,
        "mass_unit": "kg",
    },
    # -------------------------------------------------------------
    # Chair: Transport Legs
    # -------------------------------------------------------------
    {
        "process_id": "chair_materials",
        "transport_mode": "train",
        "distance": 350.0,
        "distance_unit": "km",
        "transported_mass": 14.200,
        "mass_unit": "kg",
    },
    {
        "process_id": "chair_assembly",
        "transport_mode": "truck",
        "distance": 80.0,
        "distance_unit": "km",
        "transported_mass": 14.200,
        "mass_unit": "kg",
    },
    {
        "process_id": "chair_logistics",
        "transport_mode": "truck",
        "distance": 450.0,
        "distance_unit": "km",
        "transported_mass": 14.200,
        "mass_unit": "kg",
    },
]


def style_sheet(ws, title, headers, rows):
    ws.title = title
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")  # Slate 800
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")   # Slate 50
    regular_font = Font(name="Calibri", size=10, color="0F172A")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # Write Headers
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[1].height = 26

    # Write Data Rows
    for row_idx, row_data in enumerate(rows, 2):
        is_even = (row_idx % 2 == 0)
        for col_idx, header in enumerate(headers, 1):
            val = row_data[header]
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if not is_even:
                cell.fill = zebra_fill

            # Format numbers vs strings
            if isinstance(val, (int, float)):
                if isinstance(val, float) and val.is_integer():
                    cell.number_format = "0.0"
                elif isinstance(val, float):
                    cell.number_format = "0.000" if abs(val) < 0.1 else "0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[row_idx].height = 20

    # Auto-adjust column widths
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, len(rows) + 2))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)


def create_workbook(output_filepath: str | Path):
    wb = openpyxl.Workbook()
    # remove default sheet
    wb.remove(wb.active)

    ws_prod = wb.create_sheet()
    style_sheet(
        ws_prod,
        "Products",
        ["product_id", "product_name", "functional_unit_amount", "functional_unit"],
        PRODUCTS,
    )

    ws_proc = wb.create_sheet()
    style_sheet(
        ws_proc,
        "Processes",
        ["process_id", "product_id", "process_name", "stage"],
        PROCESSES,
    )

    ws_exc = wb.create_sheet()
    style_sheet(
        ws_exc,
        "Exchanges",
        ["process_id", "flow_name", "exchange_type", "amount", "unit"],
        EXCHANGES,
    )

    ws_trans = wb.create_sheet()
    style_sheet(
        ws_trans,
        "Transportation",
        ["process_id", "transport_mode", "distance", "distance_unit", "transported_mass", "mass_unit"],
        TRANSPORTATION,
    )

    wb.save(output_filepath)
    print(f"Created styled Excel workbook at: {output_filepath}")


if __name__ == "__main__":
    create_workbook("legit_product_lca_import.xlsx")
    create_workbook("product_lca_template.xlsx")
